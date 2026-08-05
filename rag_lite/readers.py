"""Local file readers for formats that SimpleDirectoryReader would otherwise open as UTF-8 text."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, Iterable

from llama_index.core.readers.base import BaseReader
from llama_index.core.schema import Document

from rag_lite.config import AppConfig
from rag_lite.image_text import (
    ImageEnrichOptions,
    _ocr_text_looks_suspicious,
    _ocr_text_quality_summary,
    hybrid_image_to_text,
    image_enrich_options_from_config,
    merged_ingest_for_image,
)


def _extract_docx_plain_text(path: Path) -> str:
    from docx import Document as DocxDocument

    docx = DocxDocument(str(path))
    parts: list[str] = []
    for p in docx.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in docx.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _pdf_text_extract_summary(text: str) -> dict[str, int]:
    raw = str(text or "")
    stripped = raw.strip()
    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    return {
        "chars": len(stripped),
        "lines": len(lines),
    }


def _pdf_text_too_sparse(text: str, *, min_chars: int, min_lines: int) -> bool:
    stats = _pdf_text_extract_summary(text)
    return stats["chars"] < max(1, int(min_chars)) or stats["lines"] < max(1, int(min_lines))


def _merge_pdf_text_with_fallback(base_text: str, fallback_text: str) -> str:
    primary = (base_text or "").strip()
    extra = (fallback_text or "").strip()
    if not extra:
        return primary
    if not primary:
        return extra
    if extra in primary:
        return primary
    return f"{primary}\n\n{extra}"


def _csv_summary_pairs(data: dict[str, Any]) -> str:
    parts: list[str] = []
    for key, value in data.items():
        if value is None:
            continue
        if isinstance(value, bool):
            value = 1 if value else 0
        parts.append(f"{key}={value}")
    return ",".join(parts)


def _parse_summary_pairs(text: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for seg in str(text or "").split(","):
        s = seg.strip()
        if not s or "=" not in s:
            continue
        key, value = s.split("=", 1)
        key = key.strip()
        value = value.strip()
        if key:
            out[key] = value
    return out


def _summary_int(data: dict[str, Any], key: str) -> int:
    try:
        return int(data.get(key) or 0)
    except (TypeError, ValueError):
        return 0


def _pdf_scan_classification_meta(
    *,
    page_count: int,
    text_pages: int,
    text_chars: int,
    text_lines: int,
    fallback_stats: dict[str, Any] | None = None,
    fallback_text_added: bool = False,
) -> dict[str, Any]:
    stats = dict(fallback_stats or {})
    pages_rendered = _summary_int(stats, "pages_rendered")
    pages_with_text = _summary_int(stats, "pages_with_text")
    suspicious_pages = _summary_int(stats, "suspicious_pages")
    early_stopped = _summary_int(stats, "early_stopped")
    page_count = max(0, int(page_count or 0))
    text_pages = max(0, int(text_pages or 0))
    text_chars = max(0, int(text_chars or 0))
    text_lines = max(0, int(text_lines or 0))
    text_page_ratio = (text_pages / page_count) if page_count else 0.0
    suspicious_ratio = (suspicious_pages / pages_rendered) if pages_rendered else 0.0
    fallback_page_ratio = (pages_with_text / pages_rendered) if pages_rendered else 0.0
    low_text = text_chars < 500 or text_lines < 20

    if low_text and page_count <= 8 and text_pages >= 1:
        doc_class = "low_text_normal"
        reason = "few_pages_with_some_native_text"
    elif low_text and pages_rendered > 0 and suspicious_pages > 0 and suspicious_ratio >= 0.6:
        doc_class = "ocr_polluted"
        reason = "fallback_pages_mostly_suspicious"
    elif low_text and text_page_ratio <= 0.1 and pages_rendered >= min(max(page_count, 1), 8):
        if pages_with_text > 0 or fallback_text_added:
            doc_class = "scan_recoverable"
            reason = "native_text_sparse_but_fallback_recovered"
        else:
            doc_class = "scan_suspected"
            reason = "native_text_sparse_and_page_images_dominate"
    elif low_text and pages_rendered > 0 and pages_with_text == 0:
        doc_class = "extract_failed"
        reason = "fallback_attempted_without_usable_text"
    elif low_text:
        doc_class = "low_text_normal"
        reason = "low_text_without_scan_pattern"
    else:
        doc_class = "text_normal"
        reason = "native_text_sufficient"

    return {
        "pdf_doc_class": doc_class,
        "pdf_doc_class_reason": reason,
        "pdf_text_page_ratio_pct": int(round(text_page_ratio * 100)),
        "pdf_fallback_text_ratio_pct": int(round(fallback_page_ratio * 100)),
        "pdf_suspicious_page_ratio_pct": int(round(suspicious_ratio * 100)),
        "pdf_low_text": int(low_text),
        "pdf_page_fallback": int(pages_rendered > 0),
    }


def _merge_pdf_scan_meta(meta: dict[str, Any], fallback_stats: dict[str, Any] | None, *, fallback_text_added: bool) -> None:
    page_count = int(meta.get("pdf_page_count") or 0)
    text_pages = int(meta.get("pdf_text_pages") or 0)
    text_chars = int(meta.get("pdf_text_chars") or 0)
    text_lines = int(meta.get("pdf_text_lines") or 0)
    class_meta = _pdf_scan_classification_meta(
        page_count=page_count,
        text_pages=text_pages,
        text_chars=text_chars,
        text_lines=text_lines,
        fallback_stats=fallback_stats,
        fallback_text_added=fallback_text_added,
    )
    meta.update(class_meta)
    summary_data = _parse_summary_pairs(str(meta.get("image_enrichment_summary") or ""))
    summary_data.update({k: v for k, v in class_meta.items() if k in ("pdf_doc_class", "pdf_doc_class_reason")})
    meta["image_enrichment_summary"] = _csv_summary_pairs(summary_data)


def _fallback_pdf_page_images_to_text(
    path: Path,
    opts: ImageEnrichOptions,
    *,
    allow_vision: bool,
) -> tuple[str, dict[str, int]]:
    import fitz

    doc = fitz.open(str(path))
    try:
        parts: list[str] = []
        scanned_pages = 0
        pages_with_text = 0
        suspicious_pages = 0
        early_stopped = 0
        max_compact_chars = 0
        max_chinese_chars = 0
        max_latin_chars = 0
        max_digit_chars = 0
        for page_idx in range(doc.page_count):
            page = doc[page_idx]
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image_bytes = pix.tobytes("png")
            if not image_bytes:
                continue
            scanned_pages += 1
            block, _tag = hybrid_image_to_text(image_bytes, opts, allow_vision=allow_vision)
            block = (block or "").strip()
            quality = _ocr_text_quality_summary(block)
            max_compact_chars = max(max_compact_chars, int(quality.get("compact_chars") or 0))
            max_chinese_chars = max(max_chinese_chars, int(quality.get("chinese_chars") or 0))
            max_latin_chars = max(max_latin_chars, int(quality.get("latin_chars") or 0))
            max_digit_chars = max(max_digit_chars, int(quality.get("digit_chars") or 0))
            if not block:
                if scanned_pages >= 24 and pages_with_text == 0:
                    early_stopped = 1
                    break
                continue
            if bool(quality.get("suspicious")):
                suspicious_pages += 1
                if scanned_pages >= 24 and pages_with_text == 0:
                    early_stopped = 1
                    break
                continue
            pages_with_text += 1
            parts.append(f"[第{page_idx + 1}页]\n{block}")
        return "\n\n".join(parts), {
            "pages_rendered": scanned_pages,
            "pages_with_text": pages_with_text,
            "suspicious_pages": suspicious_pages,
            "early_stopped": early_stopped,
            "max_compact_chars": max_compact_chars,
            "max_chinese_chars": max_chinese_chars,
            "max_latin_chars": max_latin_chars,
            "max_digit_chars": max_digit_chars,
        }
    finally:
        doc.close()


class PlainDocxReader(BaseReader):
    """Parse .docx via python-docx (paragraphs and tables)."""

    def lazy_load_data(
        self,
        input_file: Path,
        extra_info: dict | None = None,
        **kwargs: Any,
    ) -> Iterable[Document]:
        path = Path(input_file).resolve()
        text = _extract_docx_plain_text(path)
        meta = dict(extra_info or {})
        meta.setdefault("file_name", path.name)
        meta.setdefault("file_path", str(path))
        yield Document(text=text, metadata=meta)


class PlainPdfReader(BaseReader):
    """Extract PDF text via pypdf (matches project requirements)."""

    def __init__(
        self,
        cfg: AppConfig | None = None,
        *,
        fallback_image_enrichment: bool = False,
        fallback_allow_vision: bool = False,
        pipeline_overrides: dict[str, Any] | None = None,
    ) -> None:
        self._cfg = cfg
        self._fallback_image_enrichment = bool(fallback_image_enrichment)
        self._fallback_allow_vision = bool(fallback_allow_vision)
        self._pipeline_overrides = dict(pipeline_overrides or {})

    def _fallback_opts(self) -> ImageEnrichOptions | None:
        if self._cfg is None or not self._fallback_image_enrichment:
            return None
        merged = merged_ingest_for_image(self._cfg, self._pipeline_overrides)
        merged["image_enrichment"] = True
        if not self._fallback_allow_vision:
            merged["vision_model"] = ""
        elif not str(merged.get("vision_model") or "").strip():
            merged["vision_model"] = str(self._cfg.ingest.get("vision_model") or self._cfg.ollama.get("llm_model") or "llava")
        return image_enrich_options_from_config(
            self._cfg,
            image_enrichment=True,
            pipeline_overrides=merged,
        )

    def lazy_load_data(
        self,
        input_file: Path,
        extra_info: dict | None = None,
        **kwargs: Any,
    ) -> Iterable[Document]:
        from pypdf import PdfReader

        path = Path(input_file).resolve()
        reader = PdfReader(str(path))
        texts: list[str] = []
        nonempty_pages = 0
        for page in reader.pages:
            t = page.extract_text() or ""
            t = t.strip()
            if t:
                nonempty_pages += 1
                texts.append(t)
        text = "\n\n".join(texts)
        meta = dict(extra_info or {})
        meta.setdefault("file_name", path.name)
        meta.setdefault("file_path", str(path))
        meta["pdf_text_pages"] = nonempty_pages
        meta["pdf_page_count"] = len(reader.pages)
        stats = _pdf_text_extract_summary(text)
        meta["pdf_text_chars"] = stats["chars"]
        meta["pdf_text_lines"] = stats["lines"]
        meta["image_enrichment_summary"] = ""

        fallback_opts = self._fallback_opts()
        fallback_stats: dict[str, Any] | None = None
        fallback_text_added = False
        if fallback_opts and _pdf_text_too_sparse(text, min_chars=500, min_lines=20):
            fallback_text, fallback_stats = _fallback_pdf_page_images_to_text(
                path,
                fallback_opts,
                allow_vision=self._fallback_allow_vision,
            )
            meta["image_enrichment_summary"] = _csv_summary_pairs(
                {
                    "pdf_page_fallback": 1,
                    "pages_rendered": fallback_stats.get("pages_rendered") or 0,
                    "pages_with_text": fallback_stats.get("pages_with_text") or 0,
                    "suspicious_pages": fallback_stats.get("suspicious_pages") or 0,
                    "early_stopped": fallback_stats.get("early_stopped") or 0,
                    "max_compact_chars": fallback_stats.get("max_compact_chars") or 0,
                    "max_chinese_chars": fallback_stats.get("max_chinese_chars") or 0,
                    "max_latin_chars": fallback_stats.get("max_latin_chars") or 0,
                    "max_digit_chars": fallback_stats.get("max_digit_chars") or 0,
                }
            )
            if fallback_text.strip():
                text = _merge_pdf_text_with_fallback(text, fallback_text)
                fallback_text_added = True
            elif _pdf_text_too_sparse(text, min_chars=500, min_lines=20):
                # fallback 未补充出有效内容时保留原生稀疏文本，避免静默丢失可检索信息
                meta["pdf_fallback_empty"] = 1
        _merge_pdf_scan_meta(meta, fallback_stats, fallback_text_added=fallback_text_added)
        yield Document(text=text, metadata=meta)


class HybridPdfReader(BaseReader):
    """PyMuPDF 文本 + 页内图片 OCR / 视觉模型补充。"""

    def __init__(self, opts: ImageEnrichOptions) -> None:
        self._opts = opts

    def lazy_load_data(
        self,
        input_file: Path,
        extra_info: dict | None = None,
        **kwargs: Any,
    ) -> Iterable[Document]:
        import fitz

        path = Path(input_file).resolve()
        opts = self._opts
        doc = fitz.open(str(path))
        try:
            all_parts: list[str] = []
            seen_xref: set[int] = set()
            processed = 0
            # 成功写入正文的内嵌图数量；与 raster_xrefs 区分，便于排查「页面上有图但统计为 0」
            raster_xrefs = 0
            empty_after_hybrid = 0
            for page_idx in range(doc.page_count):
                page = doc[page_idx]
                t = (page.get_text() or "").strip()
                if t:
                    all_parts.append(t)
                if processed >= opts.max_images_per_file:
                    continue
                for img_info in page.get_images(full=True) or []:
                    if processed >= opts.max_images_per_file:
                        break
                    xref = int(img_info[0])
                    if xref in seen_xref:
                        continue
                    seen_xref.add(xref)
                    try:
                        base = doc.extract_image(xref)
                    except Exception:
                        continue
                    image_bytes = base.get("image")
                    if not image_bytes:
                        continue
                    raster_xrefs += 1
                    block, tag = hybrid_image_to_text(image_bytes, opts)
                    if not block:
                        empty_after_hybrid += 1
                        continue
                    processed += 1
                    label = tag or "text"
                    all_parts.append(
                        f"\n\n[第{page_idx + 1}页 图片 {processed} · {label}]\n{block}"
                    )
            text = "\n\n".join(all_parts)
            fallback_stats: dict[str, Any] | None = None
            fallback_text_added = False
            if _pdf_text_too_sparse(text, min_chars=500, min_lines=20):
                fallback_text, fallback_stats = _fallback_pdf_page_images_to_text(
                    path,
                    opts,
                    allow_vision=bool(opts.vision_model),
                )
                if fallback_text.strip():
                    text = _merge_pdf_text_with_fallback(text, fallback_text)
                    fallback_text_added = True
                elif _pdf_text_too_sparse(text, min_chars=500, min_lines=20):
                    # fallback 未补充出有效内容时保留原生稀疏文本，避免静默丢失可检索信息
                    pass
            meta = dict(extra_info or {})
            meta.setdefault("file_name", path.name)
            meta.setdefault("file_path", str(path))
            text_stats = _pdf_text_extract_summary(text)
            native_text = "\n\n".join(part for part in all_parts if not str(part).startswith("\n\n[第"))
            native_stats = _pdf_text_extract_summary(native_text)
            native_text_pages = sum(1 for page_idx in range(doc.page_count) if (doc[page_idx].get_text() or "").strip())
            meta["pdf_text_pages"] = native_text_pages
            meta["pdf_page_count"] = doc.page_count
            meta["pdf_text_chars"] = native_stats["chars"]
            meta["pdf_text_lines"] = native_stats["lines"]
            summary_data = {
                "pdf_images": processed,
                "raster_xrefs": raster_xrefs,
                "empty_hybrid": empty_after_hybrid,
            }
            if fallback_stats is not None:
                summary_data.update(
                    {
                        "pdf_page_fallback": 1,
                        "pages_rendered": fallback_stats.get("pages_rendered") or 0,
                        "pages_with_text": fallback_stats.get("pages_with_text") or 0,
                        "suspicious_pages": fallback_stats.get("suspicious_pages") or 0,
                        "early_stopped": fallback_stats.get("early_stopped") or 0,
                        "max_compact_chars": fallback_stats.get("max_compact_chars") or 0,
                        "max_chinese_chars": fallback_stats.get("max_chinese_chars") or 0,
                        "max_latin_chars": fallback_stats.get("max_latin_chars") or 0,
                        "max_digit_chars": fallback_stats.get("max_digit_chars") or 0,
                    }
                )
            meta["image_enrichment_summary"] = _csv_summary_pairs(summary_data)
            _merge_pdf_scan_meta(meta, fallback_stats, fallback_text_added=fallback_text_added)
            meta["pdf_total_text_chars"] = text_stats["chars"]
            meta["pdf_total_text_lines"] = text_stats["lines"]
            yield Document(text=text, metadata=meta)
        finally:
            doc.close()


class HybridDocxReader(BaseReader):
    """段落/表格文本 + 压缩包内 word/media 图片（顺序与正文不完全一致）。"""

    def __init__(self, opts: ImageEnrichOptions) -> None:
        self._opts = opts

    def lazy_load_data(
        self,
        input_file: Path,
        extra_info: dict | None = None,
        **kwargs: Any,
    ) -> Iterable[Document]:
        import zipfile

        path = Path(input_file).resolve()
        opts = self._opts
        body = _extract_docx_plain_text(path)
        extra_parts: list[str] = []
        n = 0
        try:
            with zipfile.ZipFile(path) as zf:
                names = sorted(
                    n
                    for n in zf.namelist()
                    if n.startswith("word/media/") and not n.endswith("/")
                )
                for name in names:
                    if n >= opts.max_images_per_file:
                        break
                    try:
                        data = zf.read(name)
                    except Exception:
                        continue
                    if not data:
                        continue
                    block, tag = hybrid_image_to_text(data, opts)
                    if not block:
                        continue
                    n += 1
                    label = tag or "text"
                    extra_parts.append(f"\n\n[DOCX 内嵌图 {n} · {label}]\n{block}")
        except (zipfile.BadZipFile, OSError):
            pass
        text = body + "".join(extra_parts)
        meta = dict(extra_info or {})
        meta.setdefault("file_name", path.name)
        meta.setdefault("file_path", str(path))
        meta["image_enrichment_summary"] = f"docx_images={n}"
        yield Document(text=text, metadata=meta)


def default_local_file_extractors(
    cfg: AppConfig | None = None,
    *,
    image_enrichment: bool | None = None,
    pipeline_overrides: dict[str, Any] | None = None,
) -> dict[str, BaseReader]:
    """Use with SimpleDirectoryReader(file_extractor=...)."""
    if cfg is None:
        return {
            ".docx": PlainDocxReader(),
            ".pdf": PlainPdfReader(),
        }
    opts = image_enrich_options_from_config(
        cfg,
        image_enrichment=image_enrichment,
        pipeline_overrides=pipeline_overrides,
    )
    plain_pdf = PlainPdfReader(
        cfg,
        fallback_image_enrichment=True,
        fallback_allow_vision=False,
        pipeline_overrides=pipeline_overrides,
    )
    if opts is None:
        return {
            ".docx": PlainDocxReader(),
            ".pdf": plain_pdf,
        }
    return {
        ".docx": HybridDocxReader(opts),
        ".pdf": HybridPdfReader(opts),
    }
